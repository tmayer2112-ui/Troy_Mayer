#!/usr/bin/env python3
"""Build a resume in the layout of Troy's own Word resume (.docx + .pdf + .txt).

    python3 build_classic.py variants/v8_classic.json build/

Times New Roman: 18 pt bold name, 11 pt contact line, 11 pt bold caps headings with a rule,
bold employer + bold right-aligned location, bold-italic role + bold right-aligned dates,
10.5 pt bullets (bullet at 0.25 in, text at 0.5 in). Fails if the PDF is not one page.

Spec keys: name, contact (items: str or {"text","url"}), sections:
[{title, spacer, entries: [{org, org_right, role, role_link, role_right, lines: [[left, right]], bullets}]} |
 {title, rows: [[label, text]]}].
"""
import json, re, subprocess, sys
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches

from build_resume import add_runs, add_link, add_items, tight, right_tab, bottom_border

FONT, NAME, HEAD, BODY = "Times New Roman", 18, 11, 10.5
ML, MR, MT, MB = 0.555, 0.5, 0.45, 0.35
TEXT_W = 8.5 - ML - MR


def numbering(doc, left_in=0.5, hanging_in=0.25):
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum"); abstract.set(qn("w:abstractNumId"), "91")
    lvl = OxmlElement("w:lvl"); lvl.set(qn("w:ilvl"), "0")
    for tag, val in (("w:start", "1"), ("w:numFmt", "bullet"), ("w:lvlText", "•"), ("w:lvlJc", "left")):
        e = OxmlElement(tag); e.set(qn("w:val"), val); lvl.append(e)
    pPr = OxmlElement("w:pPr"); ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(int(left_in * 1440))); ind.set(qn("w:hanging"), str(int(hanging_in * 1440)))
    pPr.append(ind); lvl.append(pPr)
    abstract.append(lvl); numbering.insert(0, abstract)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), "91")
    a = OxmlElement("w:abstractNumId"); a.set(qn("w:val"), "91"); num.append(a); numbering.append(num)
    return "91"


def bullet(doc, num_id, text):
    p = doc.add_paragraph(); tight(p)
    pPr = p._p.get_or_add_pPr(); numPr = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl"); il.set(qn("w:val"), "0"); ni = OxmlElement("w:numId"); ni.set(qn("w:val"), num_id)
    numPr.append(il); numPr.append(ni); pPr.append(numPr)
    add_runs(p, text, BODY)


def spacer(doc, pts):
    p = doc.add_paragraph(); tight(p)
    r = p.add_run(""); r.font.size = Pt(pts)
    p.paragraph_format.line_spacing = Pt(pts)


def two_sided(doc, left, right, l_bold, l_italic, r_bold=True, link=None):
    p = doc.add_paragraph(); tight(p); right_tab(p, TEXT_W)
    add_runs(p, left, HEAD, bold=l_bold, italic=l_italic)
    if link:
        add_runs(p, " – ", HEAD, bold=l_bold, italic=l_italic)
        add_link(p, link["text"], link["url"], HEAD, italic=l_italic)
    if right:
        add_runs(p, "\t" + right, HEAD, bold=r_bold)


def build(spec, out_dir):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin, sec.right_margin = Inches(ML), Inches(MR)
    sec.top_margin, sec.bottom_margin = Inches(MT), Inches(MB)
    normal = doc.styles["Normal"]
    normal.font.name = FONT; normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT); normal.font.size = Pt(BODY)
    doc.core_properties.author = spec["name"]; doc.core_properties.title = f"{spec['name']} Resume"
    num_id = numbering(doc)

    p = doc.add_paragraph(); tight(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(p, spec["name"], NAME, bold=True)
    p = doc.add_paragraph(); tight(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_items(p, spec["contact"], HEAD)

    for s in spec["sections"]:
        spacer(doc, s.get("space_before", 6))
        p = doc.add_paragraph(); tight(p)
        add_runs(p, s["title"].upper(), HEAD, bold=True); bottom_border(p)
        for i, e in enumerate(s.get("entries", [])):
            if i and s.get("spacer"):
                spacer(doc, s["spacer"])
            two_sided(doc, e["org"], e.get("org_right", ""), True, False)
            if e.get("role"):
                two_sided(doc, e["role"], e.get("role_right", ""), True, True, link=e.get("role_link"))
            for left, right in e.get("lines", []):
                two_sided(doc, left, right, False, True)
            for b in e.get("bullets", []):
                bullet(doc, num_id, b)
        for label, text in s.get("rows", []):
            bullet(doc, num_id, f"**{label}:** {text}")

    out_dir.mkdir(parents=True, exist_ok=True)
    stem = spec["file_stem"]
    docx_path = out_dir / f"{stem}.docx"
    doc.save(docx_path)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
                   check=True, capture_output=True)
    pdf = out_dir / f"{stem}.pdf"
    pages = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", str(pdf)], capture_output=True,
                                                             text=True).stdout).group(1))
    subprocess.run(["pdftotext", "-layout", str(pdf), str(out_dir / f"{stem}.txt")], check=True)
    bbox = subprocess.run(["pdftotext", "-bbox", "-l", "1", str(pdf), "-"], capture_output=True, text=True).stdout
    ymax = max((float(m) for m in re.findall(r'yMax="([\d.]+)"', bbox)), default=0)
    return pages, ymax / (792 - MB * 72), pdf


if __name__ == "__main__":
    pages, fill, pdf = build(json.loads(Path(sys.argv[1]).read_text()), Path(sys.argv[2]))
    print(f"{pdf.name}: {pages} page(s), bottom of text at {fill:.0%} of usable height")
    sys.exit(0 if pages == 1 else 1)
