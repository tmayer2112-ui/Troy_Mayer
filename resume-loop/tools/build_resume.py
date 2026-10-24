#!/usr/bin/env python3
"""Build a one-page resume (.docx + .pdf + .txt) from a JSON spec.

    python3 build_resume.py variants/v0_ideal.json out/

Spec keys: name, contact (list), links (list), summary (optional str),
sections: [{title, entries: [{org, loc, role, dates, lines, bullets}]} |
           {title, rows: [[label, text]]}], style (optional overrides).
Inline **bold** is supported in bullets, lines and rows. Fails if the PDF is not one page.
"""
import json, re, subprocess, sys, shutil
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches

STYLE = {"font": "Times New Roman", "body": 10.5, "name": 17, "contact": 9.5, "head": 11,
         "margin_tb": 0.45, "margin_lr": 0.55, "bullet_indent": 0.22, "after_entry": 2.5,
         "before_section": 4.5}


UNIT = re.compile(r"(\d) (ms|s|kHz|Hz|m/s|µs|mm|in-lb|J|engineers?|days|classes|walks|years)\b")


def typeset(text):
    """Curly apostrophes, and non-breaking spaces so a number never wraps away from its unit."""
    text = re.sub(r"(\w)'(\w|\s)", r"\1’\2", text)
    return UNIT.sub("\\1\u00a0\\2", text)


def add_runs(p, text, size, bold=False, italic=False):
    text = typeset(text)
    for i, part in enumerate(re.split(r"\*\*", text)):
        if not part:
            continue
        r = p.add_run(part)
        r.font.size = Pt(size)
        r.bold = bold or (i % 2 == 1)
        r.italic = italic


def add_link(p, text, url, size, italic=False):
    """A real, clickable hyperlink run (Word and the exported PDF both keep it)."""
    r_id = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                            is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rPr.append(sz)
    if italic:
        rPr.append(OxmlElement("w:i"))
    r.append(rPr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t)
    h.append(r); p._p.append(h)


def add_items(p, items, size, sep=" | "):
    """Contact/link items: plain strings, or {"text", "url"} for a hyperlink."""
    for i, it in enumerate(items):
        if i:
            add_runs(p, sep, size)
        if isinstance(it, dict):
            add_link(p, it["text"], it["url"], size)
        else:
            add_runs(p, it, size)


def tight(p, before=0, after=0, line=1.0):
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line


def right_tab(p, width_in):
    p.paragraph_format.tab_stops.add_tab_stop(Inches(width_in), WD_TAB_ALIGNMENT.RIGHT)


def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "000000")):
        b.set(qn(f"w:{k}"), v)
    bdr.append(b)
    pPr.append(bdr)


def bullet_numbering(doc):
    """One real Word bullet list (numbering part), so ATS parsers and Word see list items."""
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "90")
    lvl = OxmlElement("w:lvl"); lvl.set(qn("w:ilvl"), "0")
    for tag, val in (("w:start", "1"), ("w:numFmt", "bullet"), ("w:lvlText", "•"), ("w:lvlJc", "left")):
        e = OxmlElement(tag); e.set(qn("w:val"), val); lvl.append(e)
    pPr = OxmlElement("w:pPr"); ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(int(STYLE["bullet_indent"] * 1440))); ind.set(qn("w:hanging"), "180")
    pPr.append(ind); lvl.append(pPr)
    abstract.append(lvl)
    numbering.insert(0, abstract)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), "90")
    a = OxmlElement("w:abstractNumId"); a.set(qn("w:val"), "90"); num.append(a)
    numbering.append(num)
    return "90"


def set_num(p, num_id):
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl"); il.set(qn("w:val"), "0")
    ni = OxmlElement("w:numId"); ni.set(qn("w:val"), num_id)
    numPr.append(il); numPr.append(ni); pPr.append(numPr)


def build(spec, out_dir):
    st = {**STYLE, **spec.get("style", {})}
    STYLE.update(st)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(st["margin_tb"])
    sec.left_margin = sec.right_margin = Inches(st["margin_lr"])
    text_w = 8.5 - 2 * st["margin_lr"]
    normal = doc.styles["Normal"]
    normal.font.name = st["font"]
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), st["font"])
    normal.font.size = Pt(st["body"])
    num_id = bullet_numbering(doc)
    doc.core_properties.author = spec["name"]  # some parsers read the name from metadata
    doc.core_properties.title = f"{spec['name']} Resume"

    p = doc.add_paragraph(); tight(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(p, spec["name"], st["name"], bold=True)
    for items in [spec["contact"]] + ([spec["links"]] if spec.get("links") else []):
        p = doc.add_paragraph(); tight(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_items(p, items, st["contact"])
    if spec.get("summary"):
        p = doc.add_paragraph(); tight(p, before=spec.get("summary_before", 4)); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, spec["summary"], st["body"], italic=spec.get("summary_italic", False))

    for s in spec["sections"]:
        p = doc.add_paragraph(); tight(p, before=st["before_section"], after=1.5)
        add_runs(p, s["title"].upper(), st["head"], bold=True); bottom_border(p)
        for e in s.get("entries", []):
            if e.get("inline"):  # one line: **org** | *role*  <tab> dates
                p = doc.add_paragraph(); tight(p, before=1); right_tab(p, text_w)
                add_runs(p, e["org"], st["body"], bold=True)
                sep = e.get("sep", " | ")
                if e.get("role"):
                    add_runs(p, sep, st["body"])
                    add_runs(p, e["role"], st["body"], italic=True)
                if e.get("link"):  # {"text", "url"} shown after the role
                    add_runs(p, sep, st["body"])
                    add_link(p, e["link"]["text"], e["link"]["url"], st["body"], italic=True)
                add_runs(p, "\t" + e.get("dates", ""), st["body"])
            elif e.get("org"):
                p = doc.add_paragraph(); tight(p, before=1); right_tab(p, text_w)
                add_runs(p, e["org"], st["body"], bold=True)
                if e.get("loc"):
                    add_runs(p, "\t" + e["loc"], st["body"], bold=e.get("loc_bold", False))
            if e.get("role") and not e.get("inline"):
                p = doc.add_paragraph(); tight(p); right_tab(p, text_w)
                add_runs(p, e["role"], st["body"], italic=True)
                if e.get("dates"):
                    add_runs(p, "\t" + e["dates"], st["body"])
            for ln in e.get("lines", []):
                p = doc.add_paragraph(); tight(p); right_tab(p, text_w)
                add_runs(p, ln, st["body"])
            for b in e.get("bullets", []):
                p = doc.add_paragraph(); tight(p); set_num(p, num_id)
                add_runs(p, b, st["body"])
            if doc.paragraphs:
                doc.paragraphs[-1].paragraph_format.space_after = Pt(st["after_entry"])
        for label, text in s.get("rows", []):
            p = doc.add_paragraph(); tight(p); set_num(p, num_id)
            add_runs(p, f"**{label}:** {text}" if label else text, st["body"])

    out_dir.mkdir(parents=True, exist_ok=True)
    stem = spec["file_stem"]
    docx_path = out_dir / f"{stem}.docx"
    doc.save(docx_path)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
                   check=True, capture_output=True)
    pdf = out_dir / f"{stem}.pdf"
    pages = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", str(pdf)], capture_output=True,
                                                             text=True).stdout).group(1))
    subprocess.run(["pdftotext", "-layout", str(pdf), str(out_dir / f"{stem}.txt")], check=True)
    # how much of the page is used: lowest text baseline vs. page height
    bbox = subprocess.run(["pdftotext", "-bbox", str(pdf), "-"], capture_output=True, text=True).stdout
    ymax = max((float(m) for m in re.findall(r'yMax="([\d.]+)"', bbox)), default=0)
    fill = ymax / (792 - st["margin_tb"] * 72)
    return pages, fill, pdf


if __name__ == "__main__":
    spec = json.loads(Path(sys.argv[1]).read_text())
    pages, fill, pdf = build(spec, Path(sys.argv[2]))
    print(f"{pdf.name}: {pages} page(s), bottom of text at {fill:.0%} of usable height")
    sys.exit(0 if pages == 1 else 1)
